"""Validate uploaded employee data and reuse transactional ERP onboarding."""
import csv
import io
from pathlib import Path
from zipfile import ZipFile, BadZipFile

from django.db import transaction
from openpyxl import load_workbook

from .erp_employee_sync import EmployeeSyncError, ingest_employee_event, validate_employee_event
from .erp_employee_mapping import normalize_external_id
from .models import Company, ERPConnector, ERPEvent


MAX_FILE_BYTES = 5 * 1024 * 1024
MAX_ROWS = 1000
MAX_SCANNED_ROWS = 10000
ALLOWED_COLUMNS = {'employee_id', 'first_name', 'name', 'last_name', 'email', 'is_active', 'mobile'}


class EmployeeImportError(Exception):
    def __init__(self, detail, *, errors=None, status=400):
        super().__init__(detail)
        self.errors = errors or []
        self.status = status


def boolean_value(value, field='is_active'):
    if isinstance(value, bool):
        return value
    if str(value).strip().lower() in ('true', '1', 'active'):
        return True
    if str(value).strip().lower() in ('false', '0', 'inactive'):
        return False
    raise ValueError(f'{field} must be true/false, 1/0 or active/inactive; do not leave it blank.')


def cell_text(value):
    if value is None:
        return ''
    if isinstance(value, str):
        return value.strip()
    if isinstance(value, int) and not isinstance(value, bool):
        return str(value)
    raise ValueError('Use text cells for IDs, names, emails and mobile numbers.')


def _xlsx_rows(content):
    # Bound expanded content before an XML reader opens the workbook.
    with ZipFile(io.BytesIO(content)) as archive:
        members = archive.infolist()
        if len(members) > 2000 or sum(m.file_size for m in members) > 20 * 1024 * 1024:
            raise EmployeeImportError('The expanded Excel file is too large. Export only employee data.')
        if any(m.flag_bits & 1 for m in members):
            raise EmployeeImportError('Password-protected Excel files are not supported.')
    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=False, keep_links=False)
    try:
        if len(workbook.worksheets) != 1:
            raise EmployeeImportError('Use one worksheet containing the employee table.')
        sheet = workbook.worksheets[0]
        sheet.reset_dimensions()
        for cells in sheet.iter_rows():
            if any(cell.data_type in ('f', 'e') for cell in cells):
                raise EmployeeImportError('Excel formulas/errors are not supported. Paste values before uploading.')
            yield [cell.value for cell in cells]
    finally:
        workbook.close()


def read_employee_file(upload):
    if upload.size > MAX_FILE_BYTES:
        raise EmployeeImportError('File must be at most 5 MiB.')
    extension = Path(upload.name).suffix.lower()
    if extension not in ('.xlsx', '.csv'):
        raise EmployeeImportError('Upload an .xlsx or UTF-8 .csv file. Save old .xls files as .xlsx first.')
    content = upload.read(MAX_FILE_BYTES + 1)
    if len(content) > MAX_FILE_BYTES:
        raise EmployeeImportError('File must be at most 5 MiB.')
    try:
        source = (csv.reader(io.StringIO(content.decode('utf-8-sig'), newline=''), strict=True)
                  if extension == '.csv' else _xlsx_rows(content))
        return _validate_rows(source)
    except EmployeeImportError:
        raise
    except (UnicodeError, csv.Error, BadZipFile):
        raise EmployeeImportError('Cannot read this file. Use a valid .xlsx or UTF-8 comma-separated CSV.')
    except Exception as exc:
        raise EmployeeImportError('Cannot read this workbook. Save a new .xlsx containing plain employee values.') from exc


def _validate_rows(source):
    rows, errors, seen_ids, seen_emails = [], [], set(), set()
    try:
        headers = [cell_text(value).lower() for value in next(source)]
    except StopIteration:
        raise EmployeeImportError('File is empty.')
    while headers and not headers[-1]:
        headers.pop()
    required = {'employee_id', 'email', 'is_active'}
    if (not required.issubset(headers) or not ({'first_name', 'name'} & set(headers)) or
            len(headers) != len(set(headers)) or '' in headers or
            set(headers) - ALLOWED_COLUMNS or {'first_name', 'name'}.issubset(headers)):
        raise EmployeeImportError('Required columns: employee_id, first_name (or name), email, is_active. '
                                  'Optional: last_name, mobile. Use one name column and unique headers.')
    count = 0
    try:
        for row_number, values in enumerate(source, start=2):
            if row_number > MAX_SCANNED_ROWS + 1:
                raise EmployeeImportError('Too many worksheet rows. Remove extra blank/formatted rows.')
            if all(value is None or value == '' for value in values):
                continue
            count += 1
            if count > MAX_ROWS:
                raise EmployeeImportError('Import at most 1000 employees per file.')
            try:
                if any(value is not None and value != '' for value in values[len(headers):]):
                    raise ValueError('Row contains values without a matching column header.')
                if len(values) < len(headers):
                    raise ValueError('Row has fewer values than the column headers.')
                data = dict(zip(headers, values))
                external_id = normalize_external_id(cell_text(data['employee_id']))
                payload = {'email': cell_text(data['email']),
                           'first_name': cell_text(data.get('first_name', data.get('name'))),
                           'is_active': boolean_value(data['is_active'])}
                if not payload['first_name']:
                    raise ValueError('Employee name cannot be blank.')
                for field in ('last_name', 'mobile'):
                    if field in data:
                        payload[field] = cell_text(data[field])
                validate_employee_event('erp_user_created', payload, external_id)
                email_key = payload['email'].casefold()
                if external_id in seen_ids:
                    raise ValueError('Duplicate employee_id in this file.')
                if email_key in seen_emails:
                    raise ValueError('Duplicate email in this file.')
                seen_ids.add(external_id)
                seen_emails.add(email_key)
                rows.append({'row': row_number, 'external_id': external_id, 'payload': payload})
            except (ValueError, EmployeeSyncError) as exc:
                errors.append({'row': row_number, 'detail': str(exc)})
    finally:
        close = getattr(source, 'close', None)
        if close:
            close()
    if errors:
        raise EmployeeImportError('File validation failed. No employees or mappings were saved.', errors=errors)
    if not rows:
        raise EmployeeImportError('File has no employee rows.')
    return rows


@transaction.atomic
def import_employee_rows(connector, token, rows, *, dry_run=False):
    Company.objects.select_for_update().get(pk=connector.company_id)
    connector = ERPConnector.objects.select_for_update().filter(
        pk=connector.pk, api_token=token, is_active=True).first()
    if connector is None:
        raise EmployeeImportError('Invalid connector token.', status=403)
    if not isinstance(connector.allowed_event_types, list) or 'erp_user_created' not in connector.allowed_event_types:
        raise EmployeeImportError('Enable erp_user_created in this connector before importing.', status=403)
    results, errors = [], []
    for row in rows:
        try:
            with transaction.atomic():
                event = ERPEvent.objects.create(company_id=connector.company_id, connector=connector,
                    event_type='erp_user_created', external_id=row['external_id'], payload=row['payload'])
                result = ingest_employee_event(event)
                result['row'] = row['row']
                if not dry_run:
                    result['event_id'] = event.pk
                elif result['action'] == 'created':
                    result.pop('user_id', None)
                    result.pop('workpulse_employee_id', None)
                results.append(result)
        except EmployeeSyncError as exc:
            errors.append({'row': row['row'], 'external_id': row['external_id'],
                           'code': exc.code, 'detail': str(exc)})
    if errors:
        raise EmployeeImportError('Import conflicts found. No employees or mappings were saved.', errors=errors, status=409)
    response = {'status': 'validated' if dry_run else 'processed', 'dry_run': dry_run,
                'company_id': connector.company_id, 'total': len(results),
                'created': sum(r['action'] == 'created' for r in results),
                'updated': sum(r['action'] == 'updated' for r in results),
                'active': sum(r['is_active'] for r in results),
                'inactive': sum(not r['is_active'] for r in results), 'employees': results}
    if dry_run:
        transaction.set_rollback(True)
    return response
