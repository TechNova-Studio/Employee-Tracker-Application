import csv
from io import BytesIO, StringIO
from unittest.mock import patch

from django.contrib.auth.models import User
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from openpyxl import Workbook
from rest_framework.authtoken.models import Token

from .models import AuditLog, Company, ERPConnector, ERPEmployeeMapping, ERPEvent, Profile, WorkSession


class ERPEmployeeFileImportTest(TestCase):
    url = '/api/erp/employees/import/'
    headers = ['employee_id', 'first_name', 'last_name', 'email', 'is_active', 'mobile']

    def setUp(self):
        self.company = Company.objects.create(name='Glutape', employee_id_prefix='WP')
        self.connector = ERPConnector.objects.create(company=self.company, api_token='file-import-token',
            allowed_event_types=['erp_user_created', 'erp_user_active', 'erp_user_inactive'])

    def row(self, index=1, active=True):
        return [f'GLUTAPE{index:04d}', 'Employee', str(index), f'employee{index}@import.test', active, '09912345678']

    def upload(self, rows=None, *, headers=None, xlsx=False, name=None):
        rows = [self.row()] if rows is None else rows
        headers = self.headers if headers is None else headers
        if xlsx:
            workbook = Workbook()
            workbook.active.append(headers)
            for row in rows:
                workbook.active.append(row)
            buffer = BytesIO()
            workbook.save(buffer)
            content = buffer.getvalue()
        else:
            buffer = StringIO(newline='')
            writer = csv.writer(buffer)
            writer.writerow(headers)
            writer.writerows(rows)
            content = buffer.getvalue().encode('utf-8-sig')
        return SimpleUploadedFile(name or ('employees.xlsx' if xlsx else 'employees.csv'), content)

    def post_file(self, upload=None, *, token=None, **fields):
        data = {'file': upload or self.upload(), **fields}
        return self.client.post(self.url, data,
            HTTP_X_WORKPULSE_TOKEN=self.connector.api_token if token is None else token)

    def existing(self, identity='GLUTAPE0001', *, company=None, email='employee1@import.test'):
        user = User.objects.create_user('existing', email=email, password='KeepExistingPassword!')
        Profile.objects.create(user=user, company=company or self.company, role='MANAGER', mobile='old')
        ERPEmployeeMapping.objects.create(connector=self.connector, external_id=identity, user=user)
        return user

    def test_xlsx_imports_148_with_exact_active_counts_and_auto_mapping(self):
        response = self.post_file(self.upload([self.row(i, i <= 27) for i in range(1, 149)], xlsx=True))
        self.assertEqual(response.status_code, 200, response.content)
        result = response.json()
        self.assertEqual((result['total'], result['created'], result['active'], result['inactive']), (148, 148, 27, 121))
        self.assertEqual(User.objects.count(), 148)
        self.assertEqual(User.objects.filter(is_active=True).count(), 27)
        self.assertEqual(Profile.objects.filter(is_active_employee=False).count(), 121)
        self.assertEqual(ERPEmployeeMapping.objects.count(), 148)
        self.assertEqual(ERPEvent.objects.filter(processed=True).count(), 148)
        mapping = ERPEmployeeMapping.objects.get(external_id='GLUTAPE0028')
        self.assertFalse(mapping.user.is_active)
        self.assertFalse(mapping.user.has_usable_password())
        self.assertEqual(mapping.user.profile.role, 'EMPLOYEE')
        self.assertEqual(mapping.user.profile.mobile, '09912345678')
        self.assertNotEqual(str(mapping.user_id), mapping.external_id)

    def test_csv_repeat_uses_same_users_and_mappings_and_preserves_role_password(self):
        user = self.existing()
        first = self.post_file(self.upload([self.row(1, False), self.row(2)]))
        self.assertEqual(first.status_code, 200, first.content)
        self.assertEqual((first.json()['created'], first.json()['updated']), (1, 1))
        ids = list(ERPEmployeeMapping.objects.values_list('external_id', 'user_id'))
        repeated = self.post_file(self.upload([self.row(1, False), self.row(2)]))
        self.assertEqual(repeated.status_code, 200, repeated.content)
        self.assertEqual((repeated.json()['created'], repeated.json()['updated']), (0, 2))
        self.assertEqual(User.objects.count(), 2)
        self.assertEqual(ids, list(ERPEmployeeMapping.objects.values_list('external_id', 'user_id')))
        user.refresh_from_db()
        self.assertTrue(user.check_password('KeepExistingPassword!'))
        self.assertEqual(user.profile.role, 'MANAGER')
        self.assertFalse(user.is_active)

    def test_dry_run_rolls_back_users_maps_events_audit_sequence_and_existing_status(self):
        user = self.existing()
        token = Token.objects.create(user=user)
        session = WorkSession.objects.create(company=self.company, user=user, status='RUNNING')
        before = AuditLog.objects.count()
        response = self.post_file(self.upload([self.row(1, False), self.row(2)]), dry_run='true')
        self.assertEqual(response.status_code, 200, response.content)
        self.assertEqual(response.json()['status'], 'validated')
        self.assertTrue(response.json()['dry_run'])
        self.assertNotIn('workpulse_employee_id', response.json()['employees'][1])
        self.assertNotIn('event_id', response.json()['employees'][0])
        self.assertEqual(User.objects.count(), 1)
        self.assertEqual(ERPEmployeeMapping.objects.count(), 1)
        self.assertEqual(ERPEvent.objects.count(), 0)
        self.assertEqual(AuditLog.objects.count(), before)
        self.company.refresh_from_db()
        self.assertEqual(self.company.next_employee_number, 1)
        user.refresh_from_db()
        session.refresh_from_db()
        self.assertTrue(user.is_active)
        self.assertEqual(session.status, 'RUNNING')
        self.assertTrue(Token.objects.filter(pk=token.pk).exists())

    def test_database_conflict_rolls_back_every_row_and_reports_sheet_row(self):
        existing = User.objects.create_user('collision', email='employee2@import.test')
        Profile.objects.create(user=existing, company=self.company)
        response = self.post_file(self.upload([self.row(1), self.row(2)]))
        self.assertEqual(response.status_code, 409, response.content)
        self.assertEqual(response.json()['errors'][0]['row'], 3)
        self.assertEqual(response.json()['errors'][0]['code'], 'mapping_required')
        self.assertEqual(User.objects.count(), 1)
        self.assertEqual(ERPEmployeeMapping.objects.count(), 0)
        self.assertEqual(ERPEvent.objects.count(), 0)
        self.assertEqual(AuditLog.objects.count(), 0)
        self.company.refresh_from_db()
        self.assertEqual(self.company.next_employee_number, 1)

    def test_conflict_rolls_back_prior_existing_employee_update_and_token_deletion(self):
        user = self.existing()
        auth = Token.objects.create(user=user)
        other = User.objects.create_user('other', email='employee2@import.test')
        Profile.objects.create(user=other, company=self.company)
        response = self.post_file(self.upload([self.row(1, False), self.row(2)]))
        self.assertEqual(response.status_code, 409)
        user.refresh_from_db()
        self.assertTrue(user.is_active)
        self.assertEqual(user.first_name, '')
        self.assertTrue(Token.objects.filter(pk=auth.pk).exists())
        self.assertEqual(ERPEvent.objects.count(), 0)

    def test_collects_invalid_rows_before_writing_and_never_defaults_missing_status(self):
        first, second = self.row(1), self.row(2)
        first[4], second[3] = '', 'bad-email'
        response = self.post_file(self.upload([first, second]))
        self.assertEqual(response.status_code, 400, response.content)
        self.assertEqual([r['row'] for r in response.json()['errors']], [2, 3])
        self.assertEqual(User.objects.count(), 0)
        self.assertEqual(ERPEvent.objects.count(), 0)

    def test_rejects_duplicate_source_ids_and_case_insensitive_emails(self):
        for changed_index in (0, 3):
            row = self.row(2)
            row[changed_index] = self.row(1)[changed_index].upper()
            response = self.post_file(self.upload([self.row(1), row]))
            self.assertEqual(response.status_code, 400, response.content)
            self.assertEqual(response.json()['errors'][0]['row'], 3)
        self.assertEqual(User.objects.count(), 0)

    def test_header_validation_missing_status_unknown_duplicate_and_ambiguous_names(self):
        for headers in (['employee_id', 'first_name', 'email'],
                        self.headers + ['company_id'], self.headers + ['email'],
                        self.headers + ['name'], ['employee_id', 'email', 'is_active']):
            response = self.post_file(self.upload(headers=headers))
            self.assertEqual(response.status_code, 400, response.content)
        self.assertEqual(User.objects.count(), 0)

    def test_name_alias_status_variants_and_numeric_source_id(self):
        response = self.post_file(self.upload([[21, 'Full Employee Name', 'person@import.test', 'INACTIVE']],
            headers=[' employee_id ', 'NAME', 'email', 'is_active'], xlsx=True))
        self.assertEqual(response.status_code, 200, response.content)
        mapping = ERPEmployeeMapping.objects.get(external_id='21')
        self.assertEqual(mapping.user.first_name, 'Full Employee Name')
        self.assertFalse(mapping.user.is_active)

    def test_connector_token_required_even_if_admin_is_logged_in(self):
        admin = User.objects.create_user('admin')
        Profile.objects.create(user=admin, company=self.company, role='COMPANY_ADMIN')
        self.client.force_login(admin)
        self.assertEqual(self.post_file(token='').status_code, 403)
        self.assertEqual(self.post_file(token='invalid').status_code, 403)
        self.connector.is_active = False
        self.connector.save()
        self.assertEqual(self.post_file().status_code, 403)
        self.assertEqual(ERPEmployeeMapping.objects.count(), 0)

    def test_allowed_event_settings_and_company_mismatch_are_rejected(self):
        self.assertEqual(self.post_file(company_id='999').status_code, 400)
        for events in (['pi_created'], 'erp_user_created', None):
            # JSONField disallows SQL NULL; use JSON object for malformed settings.
            self.connector.allowed_event_types = {} if events is None else events
            self.connector.save()
            self.assertEqual(self.post_file().status_code, 403)
        self.assertEqual(User.objects.count(), 0)

    def test_same_source_id_is_isolated_between_connector_companies(self):
        response = self.post_file()
        self.assertEqual(response.status_code, 200)
        company = Company.objects.create(name='Other')
        connector = ERPConnector.objects.create(company=company, api_token='other-token',
            allowed_event_types=['erp_user_created'])
        row = self.row()
        row[3] = 'other@import.test'
        response = self.post_file(self.upload([row]), token=connector.api_token)
        self.assertEqual(response.status_code, 200, response.content)
        self.assertEqual(response.json()['company_id'], company.pk)
        self.assertEqual(ERPEmployeeMapping.objects.filter(external_id='GLUTAPE0001').count(), 2)
        self.assertEqual(User.objects.filter(profile__company=company).count(), 1)

    def test_stale_cross_company_and_super_admin_mappings_cannot_be_imported(self):
        user = self.existing()
        for protected in ('foreign', 'super'):
            profile = user.profile
            if protected == 'foreign':
                profile.company = Company.objects.create(name='Foreign')
            else:
                profile.company = self.company
                profile.role = 'SUPER_ADMIN'
            profile.save()
            response = self.post_file()
            self.assertEqual(response.status_code, 409, response.content)
            self.assertEqual(response.json()['errors'][0]['code'], 'invalid_mapping')
            self.assertEqual(ERPEvent.objects.count(), 0)

    def test_missing_file_wrong_type_empty_file_formulas_and_multiple_sheets(self):
        self.assertEqual(self.client.post(self.url, {}, HTTP_X_WORKPULSE_TOKEN=self.connector.api_token).status_code, 400)
        self.assertEqual(self.post_file(self.upload(name='employees.xls')).status_code, 400)
        self.assertEqual(self.post_file(self.upload([])).status_code, 400)
        self.assertEqual(self.post_file(SimpleUploadedFile('broken.xlsx', b'bad')).status_code, 400)
        row = self.row()
        row[1] = '=1+1'
        self.assertEqual(self.post_file(self.upload([row], xlsx=True)).status_code, 400)
        workbook = Workbook()
        workbook.create_sheet('Other')
        buffer = BytesIO()
        workbook.save(buffer)
        self.assertEqual(self.post_file(SimpleUploadedFile('multiple.xlsx', buffer.getvalue())).status_code, 400)

    def test_json_upload_invalid_dry_run_and_extra_fields_are_rejected(self):
        response = self.client.post(self.url, {'file': 'employees.xlsx'}, content_type='application/json',
            HTTP_X_WORKPULSE_TOKEN=self.connector.api_token)
        self.assertEqual(response.status_code, 415)
        self.assertEqual(self.post_file(dry_run='maybe').status_code, 400)
        self.assertEqual(self.post_file(connector_id='1').status_code, 400)
        self.assertEqual(User.objects.count(), 0)

    def test_file_row_and_extra_cell_limits_do_not_save(self):
        response = self.post_file(SimpleUploadedFile('employees.csv', b'x' * (5 * 1024 * 1024 + 1)))
        self.assertEqual(response.status_code, 400)
        response = self.post_file(self.upload([self.row(i) for i in range(1, 1002)]))
        self.assertEqual(response.status_code, 400)
        self.assertEqual(self.post_file(self.upload([self.row() + ['extra']])).status_code, 400)
        self.assertEqual(User.objects.count(), 0)

    def test_unexpected_processing_failure_rolls_back_batch(self):
        from .erp_employee_sync import ingest_employee_event
        def processor(event):
            if event.external_id.endswith('0002'):
                raise RuntimeError('test failure')
            return ingest_employee_event(event)
        with patch('core.erp_employee_import.ingest_employee_event', side_effect=processor):
            with self.assertLogs('core.erp_employee_import_api', level='ERROR'):
                response = self.post_file(self.upload([self.row(1), self.row(2)]))
        self.assertEqual(response.status_code, 500)
        self.assertEqual(User.objects.count(), 0)
        self.assertEqual(ERPEmployeeMapping.objects.count(), 0)
        self.assertEqual(ERPEvent.objects.count(), 0)

    def test_token_rotation_after_file_validation_is_rechecked(self):
        from .erp_employee_import import read_employee_file
        def rotate(upload):
            rows = read_employee_file(upload)
            ERPConnector.objects.filter(pk=self.connector.pk).update(api_token='rotated')
            return rows
        with patch('core.erp_employee_import_api.read_employee_file', side_effect=rotate):
            response = self.post_file()
        self.assertEqual(response.status_code, 403)
        self.assertEqual(User.objects.count(), 0)
